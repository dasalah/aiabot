"""Excel report generation with openpyxl."""
import io
from typing import Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def generate_registration_report(event: dict, registrations: list[dict]) -> bytes:
    """Generate an Excel report for registrations of an event."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ثبت‌نام‌ها"
    ws.sheet_view.rightToLeft = True

    headers = [
        "ردیف", "نام و نام خانوادگی", "شماره دانشجویی", "کد ملی",
        "تلفن", "ایمیل", "نوع ثبت‌نام", "وضعیت پرداخت",
        "وضعیت تأیید", "عضو کانال", "تاریخ ثبت‌نام",
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    payment_status_map = {
        "none": "بدون پرداخت",
        "pending": "در انتظار تأیید",
        "verified": "تأیید شده",
        "rejected": "رد شده",
    }
    approval_status_map = {
        "pending": "در انتظار",
        "approved": "تأیید شده",
        "rejected": "رد شده",
        "waitlist": "لیست انتظار",
    }
    reg_type_map = {
        "free": "رایگان",
        "paid": "پرداخت شده",
        "cert": "با گواهی",
    }

    for row_idx, reg in enumerate(registrations, start=2):
        row_fill = PatternFill(
            start_color="E8F4FD" if row_idx % 2 == 0 else "FFFFFF",
            end_color="E8F4FD" if row_idx % 2 == 0 else "FFFFFF",
            fill_type="solid"
        )
        data = [
            row_idx - 1,
            reg.get("full_name", ""),
            reg.get("student_id", ""),
            reg.get("national_code", ""),
            reg.get("phone", ""),
            reg.get("email", ""),
            reg_type_map.get(reg.get("registration_type", ""), reg.get("registration_type", "")),
            payment_status_map.get(reg.get("payment_status", ""), reg.get("payment_status", "")),
            approval_status_map.get(reg.get("approval_status", ""), reg.get("approval_status", "")),
            "بله" if reg.get("channel_member") else "خیر",
            reg.get("created_at", ""),
        ]
        for col_idx, value in enumerate(data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = row_fill

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    ws.row_dimensions[1].height = 25

    # Add event info sheet
    ws_info = wb.create_sheet("اطلاعات رویداد")
    ws_info.sheet_view.rightToLeft = True
    ws_info.append(["عنوان رویداد", event.get("title", "")])
    ws_info.append(["تاریخ", event.get("event_date", "")])
    ws_info.append(["مکان", event.get("location", "")])
    ws_info.append(["ظرفیت", event.get("capacity", 0)])
    ws_info.append(["تعداد ثبت‌نام", len(registrations)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
